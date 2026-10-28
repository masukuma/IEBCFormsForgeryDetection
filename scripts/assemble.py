"""Assemble the whole affidavit (front matter, generated sections, annexure images) into one
Markdown file, a Word document and a PDF.

usage: assemble.py front.md results.json readings_merged.json gen_dir out_stem
gen_dir holds what doc_md.py, final_sections.py and annex_images.py wrote (summary.md,
stations/, s6.md, s7.md, s8.md, final_facts.json, annexures/ with index.json and kagera_top.png).

Mark-up understood (the subset the sections use): "## " and "### " headings, "^ " centred lines,
"N. " numbered paragraphs, pipe tables, ![caption](image path) images, **bold** runs.
"""
import json
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401  (A4 portrait is set by size)
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

from formguard.affidavit import _differing, exhibit_facts

front, results_path, readings_path, gen, stem = (Path(a) for a in sys.argv[1:6])
readings = json.loads(readings_path.read_text())
facts = json.loads((gen / "final_facts.json").read_text())


# --- the master markdown ------------------------------------------------------------------------
def agreement():
    """How many of the figures the findings quote both readers read alike."""
    res = json.loads(results_path.read_text())
    quoted = agree = 0
    for c in res["cases"]:
        for rep in c.get("exhibit_results", []):
            for d in _differing(exhibit_facts(c, rep, readings.get(c["id"], {}))):
                quoted += 1
                agree += d.get("model_agrees") is True
    return agree, quoted


def compact_summary():
    rows = (gen / "summary.md").read_text().strip().splitlines()[2:]
    out = ["| Polling station | PDF page | Caption | Result | How the page was made |", "|---|---|---|---|---|"]
    for r in rows:
        st, pg, cap, res, made = [c.strip() for c in r.strip("|").split("|")]
        out.append(f"| {st.replace(' polling station ', ' ')} | {re.match(r'PDF page (\d+)', pg).group(1)} | {cap} | {res} | {made} |")
    return "\n".join(out)


agree, quoted = agreement()
parts = [front.read_text().replace("{agree}", str(agree)).replace("{quoted}", str(quoted)).strip()]
parts.append("## SUMMARY OF FINDINGS\n\n16. **THAT** I examined the annexed pages for 48 polling stations. I summarise my findings in the "
             "following table and set them out station by station in the paragraphs after it.\n\n" + compact_summary())
parts.append("## FINDINGS FOR EACH POLLING STATION\n\n" + "\n\n".join(p.read_text().strip() for p in sorted((gen / "stations").glob("C*.md"))))
for name in ("s6.md", "s7.md", "s8.md"):
    parts.append((gen / name).read_text().strip())
index = json.loads((gen / "annexures" / "index.json").read_text())
annex = ["## ANNEXURES", "",
         f"Annexure [XX]-1 is the formguard software at commit {facts['commit']}, with its tests and the full numerical results of this "
         "examination, produced in electronic form.", "",
         "Each annexure that follows shows, for each figure that differs, the portal original on the left and the annexed page on the "
         "right, at the same place on the form after alignment. Below each pair are the figure, how it reads on each image and its "
         "measured similarity.", ""]
for cid, (mark, title, f) in index.items():
    annex += [f"### Annexure {mark}: {title}", "", f"![Annexure {mark}](annexures/{f})", ""]
annex += [f"### Annexure {facts['kagera_mark']}: Kagera Primary School polling station 1 of 2, the top of the sheet", "",
          f"![Annexure {facts['kagera_mark']}](annexures/kagera_top.png)", ""]
parts.append("\n".join(annex).strip())
md = "\n\n".join(parts) + "\n"
stem.with_suffix(".md").write_text(md)

# --- Word ---------------------------------------------------------------------------------------
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.54)
sec.top_margin = sec.bottom_margin = Cm(2.54)
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(6)

# page number, centred in the footer
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
    r = fp.add_run()
    if kind:
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
    else:
        el = OxmlElement("w:instrText")
        el.set(qn("xml:space"), "preserve")
        el.text = text
    r._r.append(el)


def runs(p, text, size=None):
    for k, piece in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if piece:
            r = p.add_run(piece)
            r.bold = k % 2 == 1
            if size:
                r.font.size = Pt(size)


def table(rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows if not re.match(r"^\|[-| ]+\|$", r.strip())]
    t = doc.add_table(rows=len(cells), cols=len(cells[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(cells):
        for j, text in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.size = Pt(8 if len(cells[0]) >= 4 and any(len(x) > 40 for x in row) else 9)
            r.bold = i == 0
    # columns as wide as their longest entry, within the 16 cm of text width
    longest = [max(min(len(row[j]), 60) for row in cells) + 4 for j in range(len(cells[0]))]
    t.autofit = False
    for j, n in enumerate(longest):
        for row in t.rows:
            row.cells[j].width = Cm(16.0 * n / sum(longest))
    hdr = t.rows[0]._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    hdr.append(el)  # the header row repeats on each page
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


blocks = re.split(r"\n\s*\n", md.strip())
for block in blocks:
    lines = block.splitlines()
    if lines[0].startswith("|"):
        table(lines)
        continue
    for line in lines:
        if line.startswith("## "):
            if line == "## ANNEXURES":
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            p = doc.add_paragraph()
            r = p.add_run(line[3:])
            r.bold, r.underline = True, True
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(12)
        elif line.startswith("### "):
            p = doc.add_paragraph()
            p.add_run(line[4:]).bold = True
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(6)
        elif line.startswith("^ "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            runs(p, line[2:])
        elif m := re.match(r"^!\[(.*?)\]\((.+?)\)$", line):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(gen / m.group(2)), width=Inches(6.0))
        elif m := re.match(r"^(\d+)\. (.*)$", line):
            p = doc.add_paragraph()
            # justified, unless a file name or hash would stretch the word spaces
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if re.search(r"\S{30,}", line) else WD_ALIGN_PARAGRAPH.JUSTIFY
            pf = p.paragraph_format
            pf.left_indent, pf.first_line_indent = Cm(1.27), Cm(-1.27)
            pf.tab_stops.add_tab_stop(Cm(1.27))
            p.add_run(f"{m.group(1)}.\t")
            runs(p, m.group(2))
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            runs(p, line)
docx_path = stem.with_suffix(".docx")
doc.save(docx_path)

# --- PDF ----------------------------------------------------------------------------------------
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(stem.parent), str(docx_path)],
               check=True, capture_output=True, timeout=600)
print(stem.with_suffix(".md"), docx_path, stem.with_suffix(".pdf"), f"readings agree {agree}/{quoted}")
